#!/usr/bin/env python3
# OpFocus DevOps demo deck — built ON the OpFocus branded template (10 x 5.625),
# so the template's branded layouts, corner motifs, and logo carry through.
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

NAVY=RGBColor(0x20,0x24,0x2F); SLATE=RGBColor(0x33,0x3F,0x48); BLUE=RGBColor(0x00,0x8F,0xDD)
BLUE2=RGBColor(0x3A,0x8D,0xDD); GREEN=RGBColor(0x70,0xAE,0x35); GREEN2=RGBColor(0x7A,0xC1,0x42)
INK=RGBColor(0x33,0x3F,0x48); MUTE=RGBColor(0x6B,0x76,0x82); LIGHT=RGBColor(0xF3,0xF3,0xF3)
LINE=RGBColor(0xD8,0xDE,0xE3); WHITE=RGBColor(0xFF,0xFF,0xFF); RED=RGBColor(0xC0,0x39,0x2B)
ASSET="/Users/jp/IdeaProjects/opfocus-devops/demo/assets"

prs=Presentation("/Users/jp/IdeaProjects/opfocus-devops/demo/template.pptx")
LAYOUT={l.name:l for l in prs.slide_layouts}
# clear the 12 example slides (drop relationships too, or new slides collide on part names)
_sl=prs.slides._sldIdLst
for sid in list(_sl):
    try: prs.part.drop_rel(sid.get(qn('r:id')))
    except Exception: pass
    _sl.remove(sid)

def add(layout="BLANK"):
    return prs.slides.add_slide(LAYOUT[layout])
def tb(s,text,x,y,w,h,size,color=INK,bold=False,italic=False,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,font="Arial",spacing=None):
    b=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=b.text_frame; tf.word_wrap=True
    tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0; tf.vertical_anchor=anchor
    lines=text if isinstance(text,list) else [text]
    for i,ln in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align
        r=p.add_run(); r.text=ln; f=r.font; f.size=Pt(size); f.bold=bold; f.italic=italic; f.name=font; f.color.rgb=color
    return b
def rect(s,x,y,w,h,fill=None,line=None,lw=1.0,radius=0.08,shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp=s.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h))
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb=fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb=line; sp.line.width=Pt(lw)
    sp.shadow.inherit=False
    if shape==MSO_SHAPE.ROUNDED_RECTANGLE:
        try: sp.adjustments[0]=radius
        except Exception: pass
    return sp
def oval(s,x,y,w,h,fill):
    sp=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(w),Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb=fill; sp.line.fill.background(); sp.shadow.inherit=False; return sp
def ring(s,x,y,d,color):
    sp=s.shapes.add_shape(MSO_SHAPE.DONUT,Inches(x),Inches(y),Inches(d),Inches(d))
    sp.fill.solid(); sp.fill.fore_color.rgb=color; sp.line.fill.background(); sp.shadow.inherit=False
    try: sp.adjustments[0]=0.18
    except Exception: pass
    return sp
def hline(s,x,y,w,color,wt=1.0):
    sp=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Pt(wt))
    sp.fill.solid(); sp.fill.fore_color.rgb=color; sp.line.fill.background(); sp.shadow.inherit=False; return sp
def pic(s,path,x,y,w,h):
    p=s.shapes.add_picture(path,Inches(x),Inches(y),Inches(w),Inches(h))
    p.line.color.rgb=LINE; p.line.width=Pt(0.75); return p
def center(sp): # center text vertically+both in a shape
    sp.text_frame.word_wrap=True; sp.text_frame.vertical_anchor=MSO_ANCHOR.MIDDLE
def shp_text(sp,text,size,color,bold=True,align=PP_ALIGN.CENTER,font="Arial"):
    tf=sp.text_frame; tf.word_wrap=True; tf.margin_left=Pt(2);tf.margin_right=Pt(2);tf.margin_top=Pt(1);tf.margin_bottom=Pt(1)
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE; p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=text
    f=r.font; f.size=Pt(size); f.bold=bold; f.name=font; f.color.rgb=color
def corner_motif(s):  # brand corner accents that don't block centered content
    oval(s, 8.65, 4.55, 2.4, 2.4, NAVY)      # navy arc bottom-right (mostly off-slide)
    ring(s, 9.05, -0.45, 1.25, BLUE)          # blue ring top-right
def title_text(s,text):
    tb(s,text,0.45,0.32,8.0,0.7,22,SLATE,bold=True)
    hline(s,0.47,1.02,0.9,GREEN,2.0)

# ---------- 1. TITLE (template TITLE layout) ----------
s=add("TITLE")
# fill the layout's title/subtitle placeholders
for ph in s.placeholders:
    if ph.placeholder_format.idx==0:
        ph.text=""; p=ph.text_frame.paragraphs[0]; r=p.add_run(); r.text="Salesforce DevOps, for everyone"
        r.font.size=Pt(30); r.font.bold=True; r.font.name="Arial"; r.font.color.rgb=SLATE
    elif ph.placeholder_format.idx==1:
        ph.text=""; p=ph.text_frame.paragraphs[0]; r=p.add_run()
        r.text="Promote changes and keep a clean version history — no senior developer required."
        r.font.size=Pt(13); r.font.name="Arial"; r.font.color.rgb=INK
tb(s,"CLAUDE PLUGIN  ·  BITBUCKET  ·  SALESFORCE",0.36,1.55,7,0.35,10.5,GREEN,bold=True)

# ---------- 2. PROBLEM ----------
s=add("BLANK"); corner_motif(s); title_text(s,"Why DevOps has been a weak spot")
probs=[("Change sets","Slow, manual, and they overwrite silently — the top way admins break production.",RED),
       ("The CI/CD push","Salesforce wants everyone on VCS + CI/CD, but the tooling needs a senior developer.",BLUE),
       ("Admins locked out","Non-developers can't promote cleanly, so good DevOps never reaches the team.",GREEN)]
for i,(t,d,c) in enumerate(probs):
    x=0.45+i*3.0
    rect(s,x,1.5,2.75,2.75,fill=LIGHT,line=LINE,radius=0.1)
    o=oval(s,x+0.28,1.78,0.6,0.6,c); shp_text(o,str(i+1),17,WHITE)
    tb(s,t,x+0.28,2.58,2.25,0.4,15,SLATE,bold=True)
    tb(s,d,x+0.28,3.02,2.25,1.1,11,INK)

# ---------- 3. SOLUTION ----------
s=add("BLANK"); corner_motif(s); title_text(s,"A Claude plugin that does the heavy lifting")
tb(s,"The consultant speaks in plain environment terms — “promote from Dev to Stage.” Claude translates that into the right git and Salesforce CLI steps, under OpFocus guardrails.",0.45,1.2,8.6,0.8,12.5,INK)
sol=[("Everyone can promote","Admins and devs alike — the judgment is Claude's job.",BLUE),
     ("Clean Bitbucket VCS","Every change captured to git with real history.",GREEN),
     ("Review with diffs","See exactly what changed — summary plus the line diff.",BLUE2),
     ("Shared audit trail","Every promote recorded in Salesforce.",GREEN2)]
for i,(t,d,c) in enumerate(sol):
    x=0.45+(i%2)*4.35; y=2.15+(i//2)*1.35
    rect(s,x,y,4.15,1.2,fill=WHITE,line=LINE,radius=0.1)
    rect(s,x+0.2,y+0.22,0.12,0.75,fill=c,radius=0.3)
    tb(s,t,x+0.48,y+0.16,3.5,0.4,14,SLATE,bold=True)
    tb(s,d,x+0.48,y+0.58,3.5,0.55,10.5,INK)

# ---------- 4. ARCHITECTURE ----------
s=add("BLANK"); title_text(s,"How it fits together")
def abox(x,y,w,h,t,sub,fill,tc,sc):
    r=rect(s,x,y,w,h,fill=fill,line=LINE,radius=0.1)
    tb(s,t,x+0.05,y+0.08,w-0.1,0.3,11.5,tc,bold=True,align=PP_ALIGN.CENTER)
    if sub: tb(s,sub,x+0.08,y+0.42,w-0.16,h-0.45,9,sc,align=PP_ALIGN.CENTER)
abox(3.8,1.2,2.3,0.62,"Consultant","admin or developer",LIGHT,SLATE,MUTE)
abox(3.8,2.3,2.3,0.72,"Claude plugin","promote & deploy logic",BLUE,WHITE,RGBColor(0xDC,0xEF,0xF9))
hline(s,4.95,1.82,0.02,MUTE); # tiny connector handled by arrows below
def arrow(x1,y1,x2,y2):
    c=s.shapes.add_connector(2,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    c.line.color.rgb=MUTE; c.line.width=Pt(1.25)
    le=c.line._get_or_add_ln(); he=le.makeelement(qn('a:tailEnd'),{'type':'triangle'}); le.append(he)
arrow(4.95,1.82,4.95,2.3)
abox(0.5,3.75,2.7,1.15,"Salesforce CLI","retrieve / deploy to the client orgs",WHITE,SLATE,INK)
abox(3.65,3.75,2.7,1.15,"Bitbucket Cloud","per-client repos — code & history",WHITE,SLATE,INK)
abox(6.8,3.75,2.75,1.15,"OpFocus Prod Org","audit DB + secret broker",NAVY,WHITE,RGBColor(0xC7,0xD0,0xDA))
for cx in (1.85,5.0,8.15):
    arrow(4.95,3.02,4.95,3.4); arrow(4.95,3.4,cx,3.4); arrow(cx,3.4,cx,3.75)
tb(s,"Claude orchestrates the CLI, Bitbucket, and the prod-org database — nothing new to install.",0.5,5.05,9,0.3,10,MUTE,italic=True)

# ---------- 5. DAILY FLOW ----------
s=add("BLANK"); title_text(s,"The daily flow: capture → review → promote")
flow=[("Capture","Source tracking detects exactly what changed in the org.",BLUE),
      ("Review","Every change with a plain-English summary and the line diff.",GREEN),
      ("Select","The consultant picks what to promote — nothing added behind their back.",BLUE2),
      ("Promote","Validate, deploy, commit to the branch, record it.",GREEN2)]
for i,(t,d,c) in enumerate(flow):
    y=1.3+i*0.95
    o=oval(s,0.5,y,0.68,0.68,c); shp_text(o,str(i+1),20,WHITE)
    tb(s,t,1.4,y+0.02,3.2,0.4,15,SLATE,bold=True)
    tb(s,d,1.4,y+0.4,3.3,0.52,10.5,INK)
pic(s,f"{ASSET}/04-review-ui.jpg",5.35,1.25,4.2,3.4)
tb(s,"The review screen — every change, nothing hidden",5.35,4.7,4.2,0.3,9,MUTE,italic=True,align=PP_ALIGN.CENTER)

# ---------- 6. REVIEW SCREEN ----------
s=add("BLANK"); title_text(s,"Review every change — with the real diff")
pic(s,f"{ASSET}/05-review-ui-diff.jpg",0.45,1.2,5.5,3.85)
tb(s,"Live change-selector — Apex diff expanded",0.45,5.07,5.5,0.3,9,MUTE,italic=True,align=PP_ALIGN.CENTER)
rev=[("Plain-English summary","Every row, always visible.",BLUE),
     ("Diff on demand","A chevron reveals the exact red/green line diff.",GREEN),
     ("Who & when","Last modified by and date, from the org.",BLUE2),
     ("Flags that matter","Shared-initiative and destructive deletes stand out.",GREEN2)]
for i,(t,d,c) in enumerate(rev):
    y=1.25+i*0.95
    rect(s,6.15,y,0.12,0.72,fill=c,radius=0.3)
    tb(s,t,6.42,y-0.02,3.2,0.35,12.5,SLATE,bold=True)
    tb(s,d,6.42,y+0.33,3.2,0.55,10,INK)

# ---------- 7. AUDIT TRAIL ----------
s=add("BLANK"); title_text(s,"Every promote is recorded in Salesforce")
tb(s,"The OpFocus production org is the shared database. Each promote writes a Deployment and a child record per component, each with a plain-English change summary.",0.45,1.12,9.1,0.7,12,INK)
pic(s,f"{ASSET}/01-deployment-record.jpg",0.7,1.95,2.95,3.18)
tb(s,"Deployment record",0.7,5.18,2.95,0.3,9,MUTE,italic=True,align=PP_ALIGN.CENTER)
pic(s,f"{ASSET}/02-deployment-component.jpg",3.85,1.95,2.95,3.18)
tb(s,"Component + change summary",3.85,5.18,2.95,0.3,9,MUTE,italic=True,align=PP_ALIGN.CENTER)
au=[("Who, what, where","Source/target org, branch, commit, who ran it.",BLUE),
    ("Validated → Deployed","Status lifecycle; tests captured.",GREEN),
    ("Per-component detail","Change type + a summary of what changed.",BLUE2)]
for i,(t,d,c) in enumerate(au):
    y=1.95+i*1.05
    rect(s,7.0,y,2.55,0.92,fill=LIGHT,line=LINE,radius=0.1)
    tb(s,t,7.17,y+0.1,2.25,0.35,11.5,c,bold=True)
    tb(s,d,7.17,y+0.42,2.25,0.45,9.5,INK)

# ---------- 8. SECURITY ----------
s=add("BLANK"); title_text(s,"Secrets never touch a laptop")
pic(s,f"{ASSET}/03-named-credential.jpg",0.45,1.25,4.8,3.68)
tb(s,"The token-broker Named Credential in Setup",0.45,4.96,4.8,0.3,9,MUTE,italic=True,align=PP_ALIGN.CENTER)
tb(s,"The Bitbucket OAuth secret lives only in a Named Credential in the prod org — encrypted and non-retrievable. An Apex endpoint hands out short-lived tokens. No secret file on any machine.",5.5,1.3,4.0,1.5,12,INK)
sec=[("No secret on laptops","Only disposable 2-hour tokens are vended.",BLUE),
     ("One-step offboarding","Deactivate the user — token access ends.",GREEN),
     ("Rotate in one place","Swap the secret once; everyone keeps working.",BLUE2)]
for i,(t,d,c) in enumerate(sec):
    y=2.85+i*0.8
    o=oval(s,5.5,y,0.45,0.45,c); shp_text(o,"✓",13,WHITE)
    tb(s,t,6.08,y-0.02,3.5,0.35,12,SLATE,bold=True)
    tb(s,d,6.08,y+0.32,3.5,0.4,9.5,INK)

# ---------- 9. GUARDRAILS ----------
s=add("BLANK"); title_text(s,"Guardrails built in")
g=[("Validate first, always","Check-only before any write — no surprise failures.",BLUE),
   ("Overwrite protection","Warned if the target changed under you.",GREEN),
   ("Deletes need typed confirmation","Destructive changes are off by default.",RED),
   ("Profiles excluded","Use permission sets instead — better practice.",GREEN2),
   ("Never another team's work","A missing component is flagged, never auto-added.",BLUE2),
   ("Default-namespace only","Managed packages & static resources stay out.",BLUE)]
for i,(t,d,c) in enumerate(g):
    x=0.45+(i%2)*4.6; y=1.3+(i//2)*1.22
    rect(s,x,y,4.4,1.08,fill=LIGHT,line=LINE,radius=0.1)
    o=oval(s,x+0.22,y+0.33,0.44,0.44,c); shp_text(o,"✓",14,WHITE)
    tb(s,t,x+0.82,y+0.14,3.45,0.4,12.5,SLATE,bold=True)
    tb(s,d,x+0.82,y+0.5,3.45,0.5,9.5,INK)

# ---------- 10. STATUS ----------
s=add("BLANK"); title_text(s,"Where it stands today")
tb(s,"Working end to end",0.45,1.2,4,0.4,15,GREEN,bold=True)
done=["Plugin packaged & installable from Bitbucket","Backend data model live in the sandbox",
      "Source-tracking capture + the review UI","Validated dev → stage promote (46 components, 0 errors)",
      "Token broker proven — no secret on the laptop"]
for i,t in enumerate(done):
    y=1.68+i*0.56; o=oval(s,0.5,y+0.02,0.26,0.26,GREEN); shp_text(o,"✓",10,WHITE)
    tb(s,t,0.9,y,3.9,0.45,11,INK,anchor=MSO_ANCHOR.MIDDLE)
tb(s,"Next",5.3,1.2,4,0.4,15,BLUE,bold=True)
nxt=["Roll the backend + broker to OpFocus Prod","Onboard a pilot client repo (CRL)",
     "Install for a first group of consultants","A real dev → stage deploy, then widen rollout"]
for i,t in enumerate(nxt):
    y=1.68+i*0.56; r=rect(s,5.35,y+0.02,0.26,0.26,fill=BLUE,radius=0.1); shp_text(r,str(i+1),10,WHITE)
    tb(s,t,5.75,y,3.6,0.45,11,INK,anchor=MSO_ANCHOR.MIDDLE)

# ---------- 11. CTA (template TITLE layout) ----------
s=add("TITLE")
for ph in s.placeholders:
    if ph.placeholder_format.idx==0:
        ph.text=""; p=ph.text_frame.paragraphs[0]; r=p.add_run(); r.text="Let’s promote something"
        r.font.size=Pt(28); r.font.bold=True; r.font.name="Arial"; r.font.color.rgb=SLATE
    elif ph.placeholder_format.idx==1:
        ph.text=""; p=ph.text_frame.paragraphs[0]; r=p.add_run()
        r.text="Install it from Bitbucket, point it at a sandbox, and run a real capture-and-promote."
        r.font.size=Pt(13); r.font.name="Arial"; r.font.color.rgb=INK
tb(s,"READY WHEN YOU ARE",0.36,1.55,7,0.35,10.5,GREEN,bold=True)

# ---------- 12. THANK YOU (branded full-bleed) ----------
s=add("BLANK")
s.shapes.add_picture(f"{ASSET}/brand-image2.jpg",0,0,Inches(10),Inches(5.625))

prs.save("/Users/jp/IdeaProjects/opfocus-devops/demo/OpFocus-DevOps-Demo.pptx")
print("saved", len(prs.slides._sldIdLst), "slides")
