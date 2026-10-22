#!/usr/bin/env python3
# OpFocus DevOps Technical Solution Design, built ON the branded OpFocus Word template
# (blue header band + logo + copyright footer carry through from the template).
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLUE=RGBColor(0x41,0x8F,0xDE); SLATE=RGBColor(0x33,0x3F,0x48); NAVY=RGBColor(0x20,0x24,0x2F)
GREEN=RGBColor(0x70,0xAE,0x35); INK=RGBColor(0x33,0x3F,0x48); MUTE=RGBColor(0x6B,0x76,0x82)
WHITE=RGBColor(0xFF,0xFF,0xFF); LIGHT="F3F3F3"
ASSET="/Users/jp/IdeaProjects/opfocus-devops/demo/assets"

doc=Document("/Users/jp/IdeaProjects/opfocus-devops/demo/doctemplate.docx")
# The template's single body paragraph holds the page-anchored OpFocus logo (a floating
# drawing). KEEP the first paragraph so the logo survives; drop the rest; strip the anchor
# paragraph's own spacing so content starts at the top.
_ps=list(doc.paragraphs)
for p in _ps[1:]:
    p._element.getparent().remove(p._element)
if _ps:
    _pf=_ps[0].paragraph_format; _pf.space_before=Pt(0); _pf.space_after=Pt(0); _pf.line_spacing=Pt(1)

def run(p, text, size=11, color=INK, bold=False, italic=False, font="Calibri"):
    r=p.add_run(text); r.font.size=Pt(size); r.font.color.rgb=color; r.font.bold=bold; r.font.italic=italic; r.font.name=font; return r
def para(text=None, style=None, size=11, color=INK, bold=False, italic=False, after=6, before=0, align=None):
    p=doc.add_paragraph(style=style)
    if align is not None: p.alignment=align
    p.paragraph_format.space_after=Pt(after); p.paragraph_format.space_before=Pt(before)
    if text is not None:
        if style:            # styled paragraph: inherit the template's brand style, no override
            p.add_run(text)
        else:
            run(p,text,size,color,bold,italic)
    return p
def h1(t): para(t, style="Heading 1")
def h2(t): para(t, style="Heading 2")
def bullet(text=None):
    p=doc.add_paragraph(); pf=p.paragraph_format
    pf.space_after=Pt(4); pf.left_indent=Inches(0.3); pf.first_line_indent=Inches(-0.18)
    r=p.add_run("•\t"); r.font.size=Pt(11); r.font.color.rgb=BLUE
    if text is not None: run(p,text)
    return p
def code(text):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6)
    run(p,text,size=9.5,color=NAVY,font="Consolas")
    sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:fill'),LIGHT)
    p._p.get_or_add_pPr().append(sh); return p
def shade(cell, fill):
    sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:fill'),fill)
    cell._tc.get_or_add_tcPr().append(sh)
def setw(cell,w):
    cell.width=Inches(w)
def borders(t):
    tbl=t._tbl; pr=tbl.tblPr
    b=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement(f'w:{edge}'); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4')
        e.set(qn('w:space'),'0'); e.set(qn('w:color'),'D7DCE0'); b.append(e)
    pr.append(b)
def table(headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; borders(t)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=""; setw(c,widths[i]); shade(c,"20242F")
        p=c.paragraphs[0]; run(p,h,10.5,WHITE,bold=True)
    for row in rows:
        cells=t.add_row().cells
        for i,val in enumerate(row):
            setw(cells[i],widths[i]); cells[i].text=""
            p=cells[i].paragraphs[0]; run(p,val,10,INK)
    doc.add_paragraph().paragraph_format.space_after=Pt(4)
    return t
def figure(file,w,cap):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(6); p.paragraph_format.space_after=Pt(2)
    p.add_run().add_picture(f"{ASSET}/{file}", width=Inches(w))
    c=doc.add_paragraph(); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.space_after=Pt(10)
    run(c,cap,9,MUTE,italic=True)
def toc():
    p=doc.add_paragraph(); r=p.add_run()
    for typ,txt in (("begin",None),):
        fc=OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'),typ); r._r.append(fc)
    instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text='TOC \\o "1-2" \\h \\z \\u'; r._r.append(instr)
    sep=OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'),'separate'); r._r.append(sep)
    t=OxmlElement('w:t'); t.text='Update this field (right-click → Update Field) to build the table of contents.'; r._r.append(t)
    end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end'); r._r.append(end)

# ===== Title block =====
para("TECHNICAL SOLUTION DESIGN", size=12, color=GREEN, bold=True, before=18, after=2)
para("Salesforce DevOps", style="Title")
para("A Claude plugin for capturing, promoting, and tracking Salesforce changes across client orgs with a clean Bitbucket version history.", size=13, color=INK, after=10)
para("Author: John Pipkin   ·   Date: 10/22/2026   ·   Status: Pilot (working end to end)", size=10.5, color=MUTE, after=14)
para("Contents", style="Heading 2")
_toc=["1.  Overview","2.  Goals and Non-Goals","3.  Architecture","4.  Data Model",
      "5.  Branching Model and pipeline.yml","6.  Change Detection and Capture",
      "7.  Review and Selection","8.  Dependency Handling","9.  Destructive and Mixed Deployments",
      "10. Authentication and Secret Management","11. Distribution and Installation",
      "12. Security Considerations","13. Rollout and Open Items"]
for _e in _toc:
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(2); p.paragraph_format.left_indent=Inches(0.15)
    run(p,_e,11,SLATE)
doc.add_page_break()

# 1
h1("1. Overview")
para("OpFocus manages many client Salesforce orgs. Salesforce is steering customers toward version control and CI/CD, but its native tooling realistically requires a senior developer, and change sets — the default promotion path for admins — are slow, manual, and overwrite silently. The result: good DevOps never reaches most of the team.")
para("This solution is a Claude plugin that lets every consultant — admins and developers alike — promote changes and keep a clean Bitbucket history, replacing change sets, without learning git or pre-installing specialized tooling. The consultant speaks in plain environment terms (“promote my changes from Dev to Stage”); Claude translates that into the correct git and Salesforce CLI operations under OpFocus guardrails. OpFocus’s own production org serves as the shared database that records every deployment.")
h2("Why a plugin, not a desktop app")
para("Every consultant already uses Claude, so distribution and updates are solved. The hard part of Salesforce DevOps is judgment — what to retrieve, what to promote, how to avoid overwrites — which Claude does well and a fixed GUI does not. A plugin applies that judgment under OpFocus guardrails; a cross-platform desktop app would reintroduce build, sign, and update burden for no benefit.")

# 2
h1("2. Goals and Non-Goals")
h2("Goals")
for t in ["Any consultant, including an admin who has never used git, can promote changes and keep a clean VCS.",
          "Replace the change-set workflow end to end: capture, review, promote to the next environment.",
          "Deploy to and from any org a consultant already has access to.",
          "Give every team shared visibility into what was deployed where, when, and by whom.",
          "Require no secret credentials stored on consultant laptops."]: bullet(t)
h2("Non-Goals")
for t in ["Not a git server or a replacement for Bitbucket — code lives in Bitbucket Cloud.",
          "Not Bitbucket Pipelines / CI automation (future).",
          "Not a centralized store of client-org credentials.",
          "Not a cross-platform desktop application."]: bullet(t)

# 3
h1("3. Architecture")
para("The consultant works entirely through Claude. The plugin orchestrates three systems:")
table(["Component","Role"], [
    ["Claude plugin (skills)","The orchestration layer. Detects changes, builds the review, applies guardrails, drives git and the Salesforce CLI."],
    ["Salesforce CLI (sf)","The deployment engine — retrieve and deploy metadata to/from the client orgs."],
    ["Bitbucket Cloud","Per-client repositories: the code and branch history."],
    ["OpFocus production org","The shared coordination layer: the deployment audit database and the Bitbucket secret broker."],
], [1.9,4.5])
para("The separation is deliberate: Bitbucket is the code system of record; the Salesforce org is the coordination and audit layer. No new infrastructure is introduced.")

# 4
h1("4. Data Model")
para("The backend is four custom objects deployed into the production org. It stores metadata about deployments, not the code itself.")
table(["Object","Purpose and key fields"], [
    ["Client_Repo__c","One per client repository. Bitbucket workspace / slug / URL, default branch, active flag."],
    ["Org_Connection__c","Catalog of a client’s orgs and how each maps to a git branch. Org type, alias, id, branch, promotion order, source-tracking flag."],
    ["Deployment__c","One per promote. Source/target org, branch, commit SHA, status, test level and results, deployed by/at, PR URL, approval status, notes."],
    ["Deployment_Component__c","One per component in a deployment (master-detail). Metadata type, component name, change type, and a plain-English change summary."],
], [2.0,4.4])
figure("01-deployment-record.jpg", 3.0, "Figure 1 — A Deployment record: the shared audit trail.")

# 5
h1("5. Branching Model and pipeline.yml")
para("Each client gets its own repository inside a fixed Bitbucket project. Within a repo the model is environment branches — one branch per org — and the topology varies by client. That variance lives in a per-repo file, .opfocus/pipeline.yml, so the skills read the topology instead of hard-coding it:")
code("client: Acme")
code("environments:")
code("  - { name: Dev, branch: dev, org: acme-dev }")
code("  - { name: Stage, branch: stage, org: acme-stage, validate_first: true, test_level: RunLocalTests }")
code("promotion_path: [Dev, Stage]")
para("The admin never sees a git term. “I made changes in Acme Dev” becomes a retrieve + commit to the dev branch; “promote to Stage” becomes a deploy of the selected set to the Stage org, a merge, and a recorded deployment.")

# 6
h1("6. Change Detection and Capture")
para("Capture always begins with source tracking in the client repo — the tool never hand-names components. Source tracking reports exactly what changed in the org since the last sync, replacing the manual change-set component picker.")
para("Scope: the repo’s .forceignore (written at setup from the authoritative installed-package list) excludes managed-package namespaces, static resources, and profiles. Profiles are excluded entirely — impractical in VCS/CI-CD; the tool nudges consultants to permission sets. Namespaces are never guessed from file names, which would collide with the org’s own objects.")

# 7
h1("7. Review and Selection")
para("Detected changes are presented in a locally served, SLDS-styled review screen, opened automatically in the Claude browser pane. It is a review, not just a checklist: each row shows its change type (Created / Updated / Deleted), who last modified it and when, and a plain-English summary. For updated code, a collapsible chevron reveals the exact red/green line diff; for XML metadata the summary leads. Shared-initiative warnings and destructive deletes are flagged. The consultant selects what to promote; that selection — and only that — is deployed.")
figure("05-review-ui-diff.jpg", 5.4, "Figure 2 — The review screen with an Apex diff expanded.")

# 8
h1("8. Dependency Handling")
para("The tool promotes exactly what the consultant selected. It does not auto-expand the set with whatever the target happens to be missing — a missing component may belong to another, unknown initiative, and pulling it in would promote someone else’s in-progress work.")
bullet(); run(doc.paragraphs[-1], "Structural container only. ", bold=True); run(doc.paragraphs[-1], "If a selected field, layout, or page belongs to an object not yet in the target, the object’s bare definition is included — it is the container for the work, not another initiative’s component.")
bullet(); run(doc.paragraphs[-1], "Everything else is flagged, not added. ", bold=True); run(doc.paragraphs[-1], "If a selected component references something neither selected nor in the target, the tool surfaces it for the consultant to decide.")
bullet(); run(doc.paragraphs[-1], "No validation loops. ", bold=True); run(doc.paragraphs[-1], "A single pre-deploy validation is the backstop; a missing dependency is reported as a decision, never auto-resolved.")

# 9
h1("9. Destructive and Mixed Deployments")
para("A component deleted in the source becomes a destructive change on the target, handled with more care than anything else:")
for t in ["A delete is never inferred from a missing file — the consultant explicitly confirms each deletion is intended (deletions are unchecked by default in the review).",
          "Executing any destructive change, or deleting a field/object (data loss), requires typing the component’s API name to confirm.",
          "Confirmed deletions go in destructiveChangesPost.xml alongside the additive package.xml; the mixed deploy validates check-only first."]: bullet(t)

# 10
h1("10. Authentication and Secret Management")
h2("Org authentication")
para("Each consultant authenticates the client orgs they are entitled to with the Salesforce CLI (sf org login web). Client-org credentials are never centralized.")
h2("Bitbucket access — the token broker")
para("Consultants are not Bitbucket workspace members (Atlassian bills per member), so they consume no seats. The tool authenticates to Bitbucket via a single OAuth consumer whose secret is never stored on laptops: it lives only in a Named Credential in the production org (encrypted, non-retrievable). An Apex REST endpoint makes the client-credentials call through that Named Credential and returns only a short-lived token to the authenticated, permissioned consultant.")
table(["Property","Result"], [
    ["No secret on laptops","Only disposable ~2-hour tokens are vended."],
    ["One-step offboarding","Deactivate the Salesforce user — token access ends immediately."],
    ["Central rotation","Swap the secret once in the Named Credential; every consultant keeps working."],
], [2.1,4.3])
figure("03-named-credential.jpg", 5.0, "Figure 3 — The token-broker Named Credential in Setup.")

# 11
h1("11. Distribution and Installation")
para("The plugin is distributed from a public GitHub repository via a marketplace manifest. It holds no secrets (the Bitbucket OAuth secret lives only in the prod-org Named Credential), so consultants need no repository account and nothing to authenticate. In the Claude desktop app, a consultant clicks + → Add plugins, pastes the marketplace URL, and installs opfocus-devops. The terminal / VS Code equivalent, run once in a Claude Code session, is:")
code("/plugin marketplace add https://github.com/john-pipkin-opfocus/opfocus-devops.git")
code("/plugin install opfocus-devops@opfocus-marketplace")
para("The three skills — setup-backend, setup-client-repo, and capture-and-promote — then become invokable. Active development lives in the private Bitbucket repository; GitHub is the public distribution mirror. Access to the private client code repositories is handled separately, at runtime, by the token broker — never by a per-consultant credential. (Plugins do not load in cloud sessions or under WSL; consultants use a local desktop-app, terminal, or VS Code session.)")

# 12
h1("12. Security Considerations")
for t in ["The Bitbucket OAuth secret is held only in the prod-org Named Credential, never in git, a config file, or a laptop.",
          "The token-vendor Apex endpoint is gated by the OpFocus_DevOps_User permission set.",
          "Commit attribution uses each consultant’s git identity; the Deployment record also logs who ran each promote.",
          "Deletions require typed confirmation; destructive changes to production additionally respect an approval gate.",
          "The OAuth consumer is workspace-wide write — contained by storing it only in the prod org and rotating it on a schedule."]: bullet(t)

# 13
h1("13. Rollout and Open Items")
h2("Working today")
for t in ["Plugin packaged and installable from the public GitHub marketplace.",
          "Backend data model, permission set, and token broker deployed to the DevOps sandbox and verified.",
          "Source-tracking capture and the review UI.",
          "A validated dev → stage promote (46-component closure, zero errors).",
          "Token broker proven end to end — no secret on the laptop."]: bullet(t)
h2("Next")
for t in ["Deploy the backend + broker to the OpFocus production org and enter the credential there.",
          "Onboard a pilot client repository and run a real dev → stage deploy.",
          "Install for a first group of consultants and widen the rollout.",
          "Future: Bitbucket Pipelines for CI validation; modern External Credential model."]: bullet(t)

doc.save("/Users/jp/IdeaProjects/opfocus-devops/demo/OpFocus-DevOps-Technical-Design.docx")
print("saved doc with", len(doc.paragraphs), "paragraphs,", len(doc.tables), "tables")
